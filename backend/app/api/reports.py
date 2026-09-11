import io
from datetime import datetime
from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import openpyxl

from app.database.session import get_db
from app.models.domain import User, LearnerProfile, RoleEnum, LearningGoal, AssessmentAttempt
from app.api.deps import get_current_user
from app.services.practice_analytics_service import PracticeAnalyticsService

router = APIRouter(prefix="/reports", tags=["Reports & Exports"])

def build_pdf_stream(target_user: User, db: Session) -> io.BytesIO:
    """
    Generates a professional, authoritative ReportLab PDF report for target_user
    using canonical PracticeAnalyticsService data.
    """
    summary = PracticeAnalyticsService.get_dashboard_summary(db=db, current_user=target_user)
    profile = db.query(LearnerProfile).filter(LearnerProfile.user_id == target_user.id).first()
    goals = db.query(LearningGoal).filter(LearningGoal.profile_id == profile.id).all() if profile else []

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()

    # Custom Clean Styles
    brand_style = ParagraphStyle(
        'BrandTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        spaceAfter=2
    )
    sub_title_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=10
    )
    section_head_style = ParagraphStyle(
        'SectionHead',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#0284C7'),
        spaceBefore=10,
        spaceAfter=6
    )
    cell_style = ParagraphStyle(
        'CellText',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )
    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        fontName="Helvetica-Bold",
        textColor=colors.HexColor('#0F172A')
    )

    story = []

    # 1. Header & Title Banner
    story.append(Paragraph("<b>ASL SENSEI — AI SIGN LANGUAGE PLATFORM</b>", brand_style))
    story.append(Paragraph("<b>Official Learner Performance & Alphabet Mastery Evaluation Report</b>", sub_title_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#0284C7'), spaceBefore=2, spaceAfter=8))

    # 2. Learner Information Overview
    story.append(Paragraph("<b>Learner Identification & Profile</b>", section_head_style))
    gen_time_str = datetime.utcnow().strftime("%B %d, %Y at %H:%M UTC")
    perf = summary.get("performance", {})
    
    info_data = [
        [
            Paragraph("<b>Full Name:</b>", cell_bold),
            Paragraph(target_user.full_name, cell_style),
            Paragraph("<b>Account Role:</b>", cell_bold),
            Paragraph(target_user.role.value if hasattr(target_user.role, 'value') else str(target_user.role), cell_style)
        ],
        [
            Paragraph("<b>Email / ID:</b>", cell_bold),
            Paragraph(target_user.email, cell_style),
            Paragraph("<b>Learning Level:</b>", cell_bold),
            Paragraph(summary["user"].get("learning_level", "Beginner"), cell_style)
        ],
        [
            Paragraph("<b>Preferred Language:</b>", cell_bold),
            Paragraph(summary["user"].get("preferred_language", "English"), cell_style),
            Paragraph("<b>Report Generated:</b>", cell_bold),
            Paragraph(gen_time_str, cell_style)
        ]
    ]
    t_info = Table(info_data, colWidths=[110, 160, 110, 160])
    t_info.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('PADDING', (0, 0), (-1, -1), 4),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_info)
    story.append(Spacer(1, 8))

    # 3. Lifetime Performance Metrics (100% Data-Driven)
    story.append(Paragraph("<b>Lifetime Performance Summary</b>", section_head_style))
    has_attempts = (perf.get("total_attempts") or 0) > 0
    score_display = f"{perf['overall_score']:.1f}%" if perf.get("overall_score") is not None else "No data available"
    acc_display = f"{perf['overall_accuracy']:.1f}%" if perf.get("overall_accuracy") is not None else "No data available"
    conf_display = f"{perf['average_confidence']:.1f}%" if perf.get("average_confidence") is not None else "No data available"
    streak_display = f"{perf.get('practice_streak_days', 0)} days"
    time_display = f"{perf.get('total_practice_time_mins', 0)} mins"
    sess_display = str(perf.get('total_sessions', 0))
    att_display = str(perf.get('total_attempts', 0))
    
    kpi_data = [
        [
            Paragraph("<b>Overall Score:</b>", cell_bold), Paragraph(score_display, cell_style),
            Paragraph("<b>Overall Accuracy:</b>", cell_bold), Paragraph(acc_display, cell_style)
        ],
        [
            Paragraph("<b>Avg Model Confidence:</b>", cell_bold), Paragraph(conf_display, cell_style),
            Paragraph("<b>Practice Streak:</b>", cell_bold), Paragraph(streak_display, cell_style)
        ],
        [
            Paragraph("<b>Total Practice Sessions:</b>", cell_bold), Paragraph(sess_display, cell_style),
            Paragraph("<b>Total Recorded Attempts:</b>", cell_bold), Paragraph(att_display, cell_style)
        ],
        [
            Paragraph("<b>Total Practice Time:</b>", cell_bold), Paragraph(time_display, cell_style),
            Paragraph("<b>Mastered Signs Count:</b>", cell_bold), Paragraph(f"{len(summary.get('strongest_signs', []))} / 29 signs", cell_style)
        ]
    ]
    t_kpi = Table(kpi_data, colWidths=[130, 140, 130, 140])
    t_kpi.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F1F5F9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('PADDING', (0, 0), (-1, -1), 4),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_kpi)
    story.append(Spacer(1, 8))

    # 4. Canonical 29-Sign Alphabet Mastery Matrix Breakdown
    story.append(Paragraph("<b>29-Sign Alphabet Mastery Matrix Breakdown</b>", section_head_style))
    matrix_rows = [[
        Paragraph("<b>Sign</b>", cell_bold),
        Paragraph("<b>Attempts</b>", cell_bold),
        Paragraph("<b>Correct</b>", cell_bold),
        Paragraph("<b>Accuracy</b>", cell_bold),
        Paragraph("<b>Avg Confidence</b>", cell_bold),
        Paragraph("<b>Status Classification</b>", cell_bold)
    ]]

    sign_classifications = summary.get("sign_classification", [])
    for item in sign_classifications:
        sign = item.get("sign", "")
        tot = item.get("total_attempts", 0)
        corr = item.get("correct_attempts", 0)
        acc_str = f"{item['accuracy']:.1f}%" if tot > 0 else "—"
        conf_str = f"{item['average_confidence']:.1f}%" if tot > 0 else "—"
        cat = item.get("category", "Not Attempted")
        
        # Color coding in table
        cat_color = "#047857" if cat == "Mastered" else ("#0284C7" if cat in ["Learning", "Practicing"] else ("#DC2626" if cat == "Needs Revision" else "#64748B"))
        cat_p = Paragraph(f"<b><font color='{cat_color}'>{cat}</font></b>", cell_style)
        
        matrix_rows.append([
            Paragraph(f"<b>{sign}</b>", cell_bold),
            Paragraph(str(tot), cell_style),
            Paragraph(str(corr), cell_style),
            Paragraph(acc_str, cell_style),
            Paragraph(conf_str, cell_style),
            cat_p
        ])

    t_matrix = Table(matrix_rows, colWidths=[50, 70, 70, 80, 100, 170])
    t_matrix.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F172A')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('PADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
    ]))
    story.append(t_matrix)
    story.append(Spacer(1, 10))

    # 5. Recent Actual Attempts History
    story.append(Paragraph("<b>Recent Practice Evaluation Feed</b>", section_head_style))
    recent_attempts = summary.get("recent_attempts", [])
    if recent_attempts:
        attempt_rows = [[
            Paragraph("<b>Date / Time</b>", cell_bold),
            Paragraph("<b>Target</b>", cell_bold),
            Paragraph("<b>Detected</b>", cell_bold),
            Paragraph("<b>Conf</b>", cell_bold),
            Paragraph("<b>Result</b>", cell_bold),
            Paragraph("<b>Feedback Provided</b>", cell_bold)
        ]]
        for a in recent_attempts[:10]:
            d_str = a["created_at"][:10] if a.get("created_at") else "Recent"
            res_str = "<b><font color='#047857'>CORRECT</font></b>" if a.get("is_correct") else "<b><font color='#DC2626'>INCORRECT</font></b>"
            attempt_rows.append([
                Paragraph(d_str, cell_style),
                Paragraph(f"<b>{a.get('expected_sign', '')}</b>", cell_bold),
                Paragraph(a.get('predicted_sign', ''), cell_style),
                Paragraph(f"{a.get('confidence', 0):.1f}%", cell_style),
                Paragraph(res_str, cell_style),
                Paragraph(a.get("feedback", "N/A"), cell_style)
            ])
        t_attempts = Table(attempt_rows, colWidths=[65, 45, 45, 50, 65, 270])
        t_attempts.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('PADDING', (0, 0), (-1, -1), 3),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')])
        ]))
        story.append(t_attempts)
    else:
        story.append(Paragraph("<i>No practice attempts recorded yet.</i>", cell_style))

    story.append(Spacer(1, 8))

    # 6. Learning Goals Section
    story.append(Paragraph("<b>Registered Learning Goals</b>", section_head_style))
    if goals:
        goal_rows = [[
            Paragraph("<b>Goal Description</b>", cell_bold),
            Paragraph("<b>Target Date</b>", cell_bold),
            Paragraph("<b>Status</b>", cell_bold)
        ]]
        for g in goals:
            t_str = g.target_date.strftime("%Y-%m-%d") if g.target_date else "Flexible"
            st_str = "<font color='#047857'>Completed</font>" if g.is_completed else "<font color='#0284C7'>In Progress</font>"
            goal_rows.append([
                Paragraph(g.goal_description, cell_style),
                Paragraph(t_str, cell_style),
                Paragraph(st_str, cell_style)
            ])
        t_goals = Table(goal_rows, colWidths=[300, 120, 120])
        t_goals.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('PADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t_goals)
    else:
        story.append(Paragraph("<i>No learning goals recorded.</i>", cell_style))

    story.append(Spacer(1, 10))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#94A3B8'), spaceBefore=4, spaceAfter=4))
    story.append(Paragraph(f"<i>Confidential & Authoritative Report • ASL Sensei Platform • User ID #{target_user.id}</i>", cell_style))

    doc.build(story)
    buffer.seek(0)
    return buffer

@router.get("/me/pdf")
@router.get("/performance/pdf")
@router.get("/pdf")
def generate_my_pdf_report(
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Downloads personal official PDF report for authenticated learner.
    Strictly verifies ownership (current_user.id).
    """
    buffer = build_pdf_stream(target_user=current_user, db=db)
    filename = f"ASL_Sensei_Report_{current_user.id}_{datetime.utcnow().strftime('%Y%m%d')}.pdf"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"',
        'Content-Type': 'application/pdf'
    }
    return StreamingResponse(buffer, media_type="application/pdf", headers=headers)

@router.get("/learners/{learner_id}/pdf")
def generate_learner_pdf_for_instructor_or_admin(
    learner_id: int,
    current_user: User = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Authorized drilldown PDF report download.
    Permitted ONLY for:
    - The learner themselves (current_user.id == learner_id)
    - Instructors (current_user.role == RoleEnum.INSTRUCTOR)
    - Administrators (current_user.role == RoleEnum.ADMINISTRATOR)
    """
    is_owner = (current_user.id == learner_id)
    is_instructor = (current_user.role == RoleEnum.INSTRUCTOR)
    is_admin = (current_user.role == RoleEnum.ADMINISTRATOR)

    if not (is_owner or is_instructor or is_admin):
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Unauthorized: You do not have permission to view or download this learner's report."
        )

    target_learner = db.query(User).filter(User.id == learner_id).first()
    if not target_learner:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Learner with ID {learner_id} not found."
        )

    buffer = build_pdf_stream(target_user=target_learner, db=db)
    filename = f"ASL_Sensei_Learner_{target_learner.id}_Report.pdf"
    headers = {
        'Content-Disposition': f'attachment; filename="{filename}"',
        'Content-Type': 'application/pdf'
    }
    return StreamingResponse(buffer, media_type="application/pdf", headers=headers)

@router.get("/performance")
def get_performance_report(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """JSON Performance summary endpoint using canonical analytics."""
    return PracticeAnalyticsService.get_dashboard_summary(db=db, current_user=current_user)

@router.get("/performance/excel")
@router.get("/excel")
def generate_excel_report(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    """Downloads personal Excel report for authenticated learner."""
    summary = PracticeAnalyticsService.get_dashboard_summary(db=db, current_user=current_user)
    profile = db.query(LearnerProfile).filter(LearnerProfile.user_id == current_user.id).first()
    attempts = db.query(AssessmentAttempt).filter(AssessmentAttempt.learner_id == current_user.id).order_by(AssessmentAttempt.timestamp.desc()).all()

    wb = openpyxl.Workbook()
    
    ws_profile = wb.active
    ws_profile.title = "Profile Overview"
    ws_profile.append(["Field", "Value"])
    ws_profile.append(["Full Name", current_user.full_name])
    ws_profile.append(["Email", current_user.email])
    ws_profile.append(["Role", current_user.role.value if hasattr(current_user.role, 'value') else str(current_user.role)])
    if profile:
        ws_profile.append(["Learning Level", profile.learning_level.value if hasattr(profile.learning_level, 'value') else str(profile.learning_level)])
        ws_profile.append(["Overall Performance Score (%)", profile.overall_performance_score])
        ws_profile.append(["Practice Streak (Days)", profile.practice_streak_days])
        ws_profile.append(["Total Practice Time (Mins)", profile.total_practice_time_mins])

    ws_mastery = wb.create_sheet(title="Alphabet Mastery")
    ws_mastery.append(["Sign Character", "Total Attempts", "Successful Attempts", "Average Accuracy (%)", "Average Confidence (%)", "Category"])
    for m in summary.get("sign_classification", []):
        ws_mastery.append([
            m["sign"], m["total_attempts"], m["correct_attempts"],
            m["accuracy"], m["average_confidence"], m["category"]
        ])

    ws_attempts = wb.create_sheet(title="Practice Attempts History")
    ws_attempts.append(["Attempt ID", "Expected Sign", "Predicted Sign", "Is Correct", "Confidence (%)", "Gesture Accuracy (%)", "Hand Shape Acc", "Date"])
    for a in attempts:
        ws_attempts.append([
            a.id, a.expected_sign, a.predicted_sign, a.is_correct,
            round(a.confidence * 100.0, 1) if a.confidence <= 1.0 else round(a.confidence, 1),
            a.gesture_accuracy, a.hand_shape_accuracy, str(a.timestamp)
        ])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    headers = {'Content-Disposition': f'attachment; filename="sign_language_report_{current_user.id}.xlsx"'}
    return StreamingResponse(buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers=headers)
